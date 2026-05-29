import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'c:\Users\최다빈\Desktop\메모패드_제작업체비교_최종.pptx')
print(f'슬라이드 수: {len(prs.slides)}')
print(f'크기: {prs.slide_width.inches:.2f}" x {prs.slide_height.inches:.2f}"')

for i, slide in enumerate(prs.slides):
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t)
        elif sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        texts.append(t)
    print(f'--- 슬라이드 {i+1} ---')
    for t in texts[:20]:
        print(' ', t)
