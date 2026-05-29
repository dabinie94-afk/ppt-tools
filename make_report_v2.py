# -*- coding: utf-8 -*-
"""
CEO 보고용 PPT - 신규입사자 지원물품 지급기준 및 개선방안
GA Intelligence | 총무팀 | 2026.05
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ── 색상 ─────────────────────────────────────────────────
NAVY        = RGBColor(0x0D, 0x2B, 0x55)
BLUE        = RGBColor(0x1A, 0x5F, 0xA8)
BLUE_LIGHT  = RGBColor(0xE8, 0xF2, 0xFC)
TEAL        = RGBColor(0x00, 0x7A, 0x87)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_BG     = RGBColor(0xF5, 0xF7, 0xFA)
GRAY_DARK   = RGBColor(0x2E, 0x2E, 0x2E)
GRAY_MID    = RGBColor(0x6B, 0x7C, 0x93)
GRAY_LINE   = RGBColor(0xD8, 0xDE, 0xE9)
RED         = RGBColor(0xC0, 0x39, 0x2B)
RED_LIGHT   = RGBColor(0xFD, 0xED, 0xEC)
GREEN       = RGBColor(0x1A, 0x7A, 0x4A)
GREEN_LIGHT = RGBColor(0xE8, 0xF8, 0xEE)
ORANGE      = RGBColor(0xD4, 0x6E, 0x0A)
ORANGE_LIGHT= RGBColor(0xFE, 0xF3, 0xE2)
GOLD        = RGBColor(0xF5, 0xA6, 0x23)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank = prs.slide_layouts[6]


# ════════════════════════════════════════
# 유틸리티
# ════════════════════════════════════════
def rgb_hex(rgb):
    return '{:02X}{:02X}{:02X}'.format(rgb[0], rgb[1], rgb[2])

def rect(slide, l, t, w, h, fill=None, line=None, lw=Pt(0), radius=0):
    shp = slide.shapes.add_shape(1, l, t, w, h)
    shp.line.width = lw
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = line
    else:
        shp.line.fill.background()
    return shp

def txbox(slide, text, l, t, w, h,
          size=Pt(11), bold=False, color=GRAY_DARK,
          align=PP_ALIGN.LEFT, wrap=True, italic=False,
          line_spacing=None):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size  = size
    run.font.bold  = bold
    run.font.color.rgb = color
    run.font.name  = "맑은 고딕"
    run.font.italic = italic
    return tb

def cell_bg(cell, rgb):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    # 기존 fill 제거
    for old in tcPr.findall(qn('a:solidFill')):
        tcPr.remove(old)
    sf = etree.SubElement(tcPr, qn('a:solidFill'))
    sr = etree.SubElement(sf,   qn('a:srgbClr'))
    sr.set('val', rgb_hex(rgb))

def cell_border(cell, color=GRAY_LINE, width=Pt(0.5)):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for side in ('a:lnL','a:lnR','a:lnT','a:lnB'):
        ln = etree.SubElement(tcPr, qn(side))
        ln.set('w', str(int(width)))
        sf = etree.SubElement(ln, qn('a:solidFill'))
        sr = etree.SubElement(sf, qn('a:srgbClr'))
        sr.set('val', rgb_hex(color))

def stylecell(cell, text, size=Pt(10), bold=False,
              color=GRAY_DARK, align=PP_ALIGN.CENTER, bg=None,
              border=True, wrap=True, italic=False):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name   = "맑은 고딕"
    run.font.size   = size
    run.font.bold   = bold
    run.font.color.rgb = color
    run.font.italic = italic
    if bg:
        cell_bg(cell, bg)
    if border:
        cell_border(cell, GRAY_LINE, Pt(0.5))

def header_bar(slide, title, page):
    """공통 상단 헤더"""
    rect(slide, 0, 0, W, Inches(0.72), fill=NAVY)
    # 왼쪽 포인트 라인
    rect(slide, 0, 0, Inches(0.06), Inches(0.72), fill=GOLD)
    txbox(slide, title,
          Inches(0.2), Inches(0.1), Inches(10), Inches(0.55),
          size=Pt(18), bold=True, color=WHITE)
    txbox(slide, page,
          Inches(12.0), Inches(0.18), Inches(1.2), Inches(0.38),
          size=Pt(10), color=RGBColor(0xAA, 0xBB, 0xCC), align=PP_ALIGN.RIGHT)
    # 하단 푸터
    rect(slide, 0, H - Inches(0.32), W, Inches(0.32), fill=NAVY)
    txbox(slide, "GA Intelligence  ·  총무팀  ·  대외비  ·  2026. 05",
          Inches(0.3), H - Inches(0.3), Inches(8), Inches(0.28),
          size=Pt(8), color=RGBColor(0xAA, 0xBB, 0xCC))

def kpi_card(slide, l, t, w, h, num, label, desc, accent):
    """KPI 카드 컴포넌트"""
    rect(slide, l, t, w, h, fill=WHITE, line=GRAY_LINE, lw=Pt(0.8))
    rect(slide, l, t, Inches(0.07), h, fill=accent)
    txbox(slide, num,
          l + Inches(0.18), t + Inches(0.1), w - Inches(0.25), Inches(0.6),
          size=Pt(28), bold=True, color=accent, align=PP_ALIGN.LEFT)
    txbox(slide, label,
          l + Inches(0.18), t + Inches(0.65), w - Inches(0.25), Inches(0.4),
          size=Pt(12), bold=True, color=NAVY, align=PP_ALIGN.LEFT)
    txbox(slide, desc,
          l + Inches(0.18), t + Inches(1.1), w - Inches(0.3), h - Inches(1.2),
          size=Pt(10), color=GRAY_MID, align=PP_ALIGN.LEFT, wrap=True)


# ════════════════════════════════════════
# SLIDE 1 — 표지
# ════════════════════════════════════════
slide = prs.slides.add_slide(blank)

# 배경 분할 (좌: 네이비, 우: 흰색)
rect(slide, 0, 0, Inches(5.8), H, fill=NAVY)
rect(slide, Inches(5.8), 0, W - Inches(5.8), H, fill=WHITE)

# 좌측 골드 포인트 라인
rect(slide, Inches(0.6), Inches(1.6), Inches(0.06), Inches(4.0), fill=GOLD)

# 메인 제목
txbox(slide, "신규입사자 지원물품\n지급기준 및 개선방안",
      Inches(0.8), Inches(1.7), Inches(4.8), Inches(2.2),
      size=Pt(34), bold=True, color=WHITE)

# 부제
txbox(slide, "사원증 재발급 기준 현실화를 통한\n비용 절감 및 보안 강화",
      Inches(0.8), Inches(4.0), Inches(4.8), Inches(1.0),
      size=Pt(13), color=RGBColor(0xB0, 0xC8, 0xE8))

# 구분선
rect(slide, Inches(0.8), Inches(5.2), Inches(4.0), Inches(0.03), fill=GOLD)

# 보고 정보
info = [
    ("보 고 부 서", "GA Intelligence  |  총무팀"),
    ("보 고 일 자", "2026년 5월"),
    ("문 서 등 급", "대  외  비"),
]
for i, (label, val) in enumerate(info):
    ty = Inches(5.35) + i * Inches(0.42)
    txbox(slide, label,
          Inches(0.8), ty, Inches(1.5), Inches(0.38),
          size=Pt(9.5), color=GOLD)
    txbox(slide, val,
          Inches(2.35), ty, Inches(3.0), Inches(0.38),
          size=Pt(9.5), color=RGBColor(0xCC, 0xDD, 0xEE))

# 우측 — 보고서 요약 카드
rect(slide, Inches(6.2), Inches(1.0), Inches(6.7), Inches(5.8),
     fill=GRAY_BG, line=GRAY_LINE, lw=Pt(0.8))
txbox(slide, "보고서 핵심 개요",
      Inches(6.5), Inches(1.2), Inches(6.0), Inches(0.45),
      size=Pt(14), bold=True, color=NAVY)
rect(slide, Inches(6.5), Inches(1.67), Inches(5.8), Inches(0.03), fill=BLUE)

summary_items = [
    ("■  배경",    "사원증 무상 재발급 남용 사례 증가 → 비용 누수 및 보안 위협"),
    ("■  제안",    "분실·파손 시 급여 공제 방식으로 전환 (3종 한정)"),
    ("■  대상",    "사원증 · 사원증 케이스 · 사원증 목걸이 (총 3종)"),
    ("■  예외",    "근속 5년 이상 자연 노후화 → 1회 무상 교체 허용"),
    ("■  효과",    "재경 비용 절감 + 보안 강화 + 직원 불만 최소화"),
    ("■  일정",    "2026.06.01 변경 기준 전격 시행"),
]
for i, (k, v) in enumerate(summary_items):
    ty = Inches(1.85) + i * Inches(0.72)
    txbox(slide, k,
          Inches(6.5), ty, Inches(1.2), Inches(0.38),
          size=Pt(10.5), bold=True, color=BLUE)
    txbox(slide, v,
          Inches(7.6), ty, Inches(5.1), Inches(0.58),
          size=Pt(10.5), color=GRAY_DARK, wrap=True)


# ════════════════════════════════════════
# SLIDE 2 — Executive Summary (핵심 요약)
# ════════════════════════════════════════
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, W, H, fill=GRAY_BG)
header_bar(slide, "Executive Summary  ·  핵심 요약", "2 / 7")

# 중앙 타이틀
rect(slide, Inches(0.4), Inches(0.85), W - Inches(0.8), Inches(0.52), fill=BLUE_LIGHT)
txbox(slide, "이 보고서는 사원증 재발급 제도 개선을 통한 비용 절감 및 출입 보안 강화 방안에 대한 CEO 결재를 요청합니다.",
      Inches(0.55), Inches(0.9), W - Inches(1.1), Inches(0.42),
      size=Pt(11.5), color=NAVY, bold=False, italic=True)

# 4개 박스: 현황 문제 / 개선안 / 기대효과 / 요청사항
boxes = [
    (NAVY,   "🔴  현황 및 문제",
     "현행 사원증 재발급 기준\n"
     "• 분실·파손 시 횟수 무제한 무상 지급\n"
     "• 사원증 세트(케이스·목걸이 포함) 동일 적용\n"
     "• 분실 건수 증가 → 예산 누수 및 보안 취약"),
    (RED,    "🟡  개선(안) 핵심",
     "변경 대상 3종 재발급 기준 강화\n"
     "• 분실·파손 시 급여 100% 공제 후 재지급\n"
     "• 근속 5년↑ 자연 노후화 → 1회 무상 예외\n"
     "• 최초 지급 및 나머지 11개 품목 현행 유지"),
    (GREEN,  "🟢  기대 효과",
     "3가지 핵심 성과 기대\n"
     "• 연간 사원증 제작·부속품 비용 절감\n"
     "• 분실 경각심 고취 → 출입 보안 사고 예방\n"
     "• 합리적 예외 조항으로 직원 반발 최소화"),
    (ORANGE, "📋  결재 요청",
     "CEO 승인 후 즉시 추진 가능\n"
     "• 2026.05   사내 규정 개정 / 인사·재경팀 연동\n"
     "• 2026.05   전사 공지 (유예기간 1주일)\n"
     "• 2026.06.01  변경 기준 전격 시행"),
]

for i, (accent, title, body) in enumerate(boxes):
    col = i % 2
    row = i // 2
    lft = Inches(0.4) + col * Inches(6.45)
    top = Inches(1.55) + row * Inches(2.65)
    bw  = Inches(6.1)
    bh  = Inches(2.45)
    rect(slide, lft, top, bw, bh, fill=WHITE, line=GRAY_LINE, lw=Pt(0.8))
    rect(slide, lft, top, bw, Inches(0.5), fill=accent)
    txbox(slide, title,
          lft + Inches(0.15), top + Inches(0.06), bw - Inches(0.3), Inches(0.4),
          size=Pt(12), bold=True, color=WHITE)
    txbox(slide, body,
          lft + Inches(0.2), top + Inches(0.6), bw - Inches(0.4), bh - Inches(0.75),
          size=Pt(10.5), color=GRAY_DARK, wrap=True)


# ════════════════════════════════════════
# SLIDE 3 — 추진 배경 및 목적
# ════════════════════════════════════════
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, W, H, fill=GRAY_BG)
header_bar(slide, "01  추진 배경 및 목적", "3 / 7")

# 현황 박스 (상단)
rect(slide, Inches(0.4), Inches(0.88), W - Inches(0.8), Inches(1.5),
     fill=RED_LIGHT, line=RED, lw=Pt(1.0))
txbox(slide, "⚠  현황 : 사원증 재발급 제도 문제점",
      Inches(0.6), Inches(0.95), Inches(8), Inches(0.38),
      size=Pt(13), bold=True, color=RED)
txbox(slide,
      "사원증 및 부속품(케이스·목걸이)은 현재 분실·파손 시 횟수 제한 없이 무상으로 재지급되고 있으며, "
      "이를 악용한 부주의 구제 사례가 지속적으로 증가하고 있습니다.",
      Inches(0.6), Inches(1.38), W - Inches(1.2), Inches(0.85),
      size=Pt(11), color=GRAY_DARK, wrap=True)

# 2가지 추진 이유
reasons = [
    (NAVY, "자산 관리 책임감 부여",
     "현행",
     "횟수 무제한 무상 재지급",
     "문제",
     "부주의 분실 구제 사례 반복 증가\n직원의 자산 관리 의식 저하\n출입 보안 위협 증가",
     "개선",
     "분실·파손 → 급여 공제 후 재지급\n자연 노후화(5년↑) 예외 조항 신설"),
    (BLUE, "비용 절감 및 행정 효율화",
     "비용",
     "사원증: 8,250원 / 재발급: 46,750원\n케이스: 30,800원  목걸이: 7,700원",
     "문제",
     "무분별 재발급으로 연간 비용 누수\n총무팀 행정 처리 공수 반복 발생",
     "효과",
     "소모성 재경 비용 절감\n총무팀 행정 효율화 기대"),
]

for i, (accent, title, l1, v1, l2, v2, l3, v3) in enumerate(reasons):
    lft = Inches(0.4) + i * Inches(6.45)
    top = Inches(2.55)
    bw  = Inches(6.1)
    bh  = Inches(4.5)
    rect(slide, lft, top, bw, bh, fill=WHITE, line=GRAY_LINE, lw=Pt(0.8))
    rect(slide, lft, top, bw, Inches(0.52), fill=accent)
    txbox(slide, f"{'①' if i==0 else '②'}  {title}",
          lft + Inches(0.18), top + Inches(0.07), bw - Inches(0.3), Inches(0.4),
          size=Pt(13), bold=True, color=WHITE)
    # 내부 3개 소박스
    for j, (lbl, val, bg_c) in enumerate([
        (l1, v1, BLUE_LIGHT),
        (l2, v2, RED_LIGHT),
        (l3, v3, GREEN_LIGHT),
    ]):
        sy = top + Inches(0.68) + j * Inches(1.2)
        rect(slide, lft + Inches(0.15), sy, bw - Inches(0.3), Inches(1.1),
             fill=bg_c, line=GRAY_LINE, lw=Pt(0.5))
        txbox(slide, lbl,
              lft + Inches(0.25), sy + Inches(0.06), Inches(1.0), Inches(0.28),
              size=Pt(9), bold=True,
              color=NAVY if bg_c==BLUE_LIGHT else (RED if bg_c==RED_LIGHT else GREEN))
        txbox(slide, val,
              lft + Inches(0.25), sy + Inches(0.35), bw - Inches(0.5), Inches(0.7),
              size=Pt(10.5), color=GRAY_DARK, wrap=True)


# ════════════════════════════════════════
# SLIDE 4 — 현행 vs 변경(안) 핵심 비교
# ════════════════════════════════════════
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, W, H, fill=GRAY_BG)
header_bar(slide, "02  현행 vs 변경(안) 비교  ·  조정 대상 3종", "4 / 7")

# 상단 안내
rect(slide, Inches(0.4), Inches(0.85), W - Inches(0.8), Inches(0.42),
     fill=BLUE_LIGHT, line=BLUE, lw=Pt(0.8))
txbox(slide,
      "조정 대상 : 사원증 · 사원증 케이스 · 사원증 목걸이 (총 3종)  |  나머지 11개 품목은 현행 기준 엄격 유지",
      Inches(0.6), Inches(0.88), W - Inches(1.2), Inches(0.35),
      size=Pt(11.5), bold=True, color=NAVY)

# 비교 테이블
compare = [
    ("구  분",       "현  행  (AS-IS)",                        "변경(안)  (TO-BE)",              "단  가"),
    ("사원증",       "분실·파손 시\n횟수 제한 없이 무상 재지급", "분실·파손 시\n급여 100% 공제 후 재지급",  "신규: 8,250원\n재발급: 46,750원"),
    ("사원증 케이스","좌동",                                    "좌동 → 급여 공제",               "30,800원"),
    ("사원증 목걸이","좌동",                                    "좌동 → 급여 공제",               "7,700원"),
    ("예외 조항",    "(없음)",                                  "근속 5년 이상 자연 노후화 시\n1회 무상 교체 허용", "인사팀 확인 후 처리"),
]

col_ws = [Inches(1.9), Inches(3.6), Inches(4.4), Inches(2.75)]
tbl = slide.shapes.add_table(
    len(compare), 4,
    Inches(0.4), Inches(1.4),
    sum(col_ws), Inches(4.9)
).table
for ci, cw in enumerate(col_ws):
    tbl.columns[ci].width = cw

for ri, row in enumerate(compare):
    for ci, val in enumerate(row):
        c = tbl.cell(ri, ci)
        if ri == 0:
            stylecell(c, val, size=Pt(11.5), bold=True, color=WHITE, bg=NAVY)
        else:
            bg = WHITE if ri % 2 == 1 else BLUE_LIGHT
            if ci == 0:
                stylecell(c, val, size=Pt(12), bold=True, color=NAVY, bg=bg)
            elif ci == 1:
                # 현행: 붉은 강조
                stylecell(c, val, size=Pt(11), color=RED, bg=RED_LIGHT if ri in (1,2,3) else bg,
                          bold=(ri in (1,2,3)))
            elif ci == 2:
                # 변경안: 초록 강조
                stylecell(c, val, size=Pt(11), color=GREEN, bg=GREEN_LIGHT if ri in (1,2,3) else bg,
                          bold=(ri in (1,2,3)))
            else:
                stylecell(c, val, size=Pt(10.5), color=GRAY_MID, bg=bg)

# 범례
rect(slide, Inches(0.4), Inches(6.48), Inches(5.5), Inches(0.3),
     fill=RED_LIGHT, line=GRAY_LINE, lw=Pt(0.5))
txbox(slide, "■  현행 (AS-IS) : 변경 전 무상 재지급 방식",
      Inches(0.6), Inches(6.5), Inches(5.0), Inches(0.26),
      size=Pt(9.5), color=RED)
rect(slide, Inches(6.1), Inches(6.48), Inches(5.5), Inches(0.3),
     fill=GREEN_LIGHT, line=GRAY_LINE, lw=Pt(0.5))
txbox(slide, "■  변경(안) (TO-BE) : 급여 공제 방식으로 전환",
      Inches(6.3), Inches(6.5), Inches(5.0), Inches(0.26),
      size=Pt(9.5), color=GREEN)


# ════════════════════════════════════════
# SLIDE 5 — 품목별 지급기준 전체표
# ════════════════════════════════════════
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, W, H, fill=GRAY_BG)
header_bar(slide, "03  품목별 지급기준 전체표  ·  신규입사자 지원물품 14종", "5 / 7")

full = [
    # 품목, 최초지급, 현행 재지급, 변경 재지급, 예산부서, 단가, 변경여부
    ("사원증",         "입사 시 1매",          "횟수 무제한 무상",     "급여 공제 후 재지급",  "총무팀",    "8,250",      True),
    ("사원증 케이스",  "입사 시 1개",          "횟수 무제한 무상",     "급여 공제 후 재지급",  "총무팀",    "30,800",     True),
    ("사원증 목걸이",  "입사 시 1개",          "횟수 무제한 무상",     "급여 공제 후 재지급",  "총무팀",    "7,700",      True),
    ("검사화",         "입사 시 1켤레",        "요청 시 발주·지급",    "좌동",                 "각 부서",   "49,500~\n151,800", False),
    ("검사가운",       "입사 시 2벌",          "요청 시 발주·지급",    "좌동",                 "각 부서",   "27,500",     False),
    ("검사복",         "입사 시 1벌",          "요청 시 발주·지급",    "좌동",                 "각 부서",   "74,800",     False),
    ("전문의가운",     "입사 시 2벌",          "요청 시 발주·지급",    "좌동",                 "각 부서",   "77,000",     False),
    ("동계피복(점퍼)", "입사 연도 최초",       "없음",                 "좌동",                 "총무팀",    "100,000",    False),
    ("동계피복(조끼)", "입사 연도 최초",       "없음",                 "좌동",                 "총무팀",    "90,300",     False),
    ("명함",           "오피스디포 신청",      "필요 시 개별 신청",    "좌동",                 "—",         "10,309",     False),
    ("다이어리·달력",  "입사 시 1권·1부",      "없음",                 "좌동",                 "학술홍보팀","—",          False),
    ("법인폰 (지점)",  "입사 시 1개",          "없음",                 "좌동",                 "총무팀",    "30,000",     False),
    ("법인차량 (임원·전문의)", "직급별 기준",  "—",                    "좌동",                 "—",         "직급별 상이",False),
    ("태블릿 (임원·전문의)",   "입사 시 1개",  "없음",                 "좌동",                 "—",         "1,200,000",  False),
]

headers = ["품목명", "최초 지급", "현행 재지급", "변경 재지급(안)", "예산부서", "단가(원)"]
col_ws2 = [Inches(1.85), Inches(1.55), Inches(2.15), Inches(2.45), Inches(1.2), Inches(1.4)]

tbl2 = slide.shapes.add_table(
    len(full)+1, 6,
    Inches(0.35), Inches(0.88),
    sum(col_ws2), Inches(6.28)
).table
for ci, cw in enumerate(col_ws2):
    tbl2.columns[ci].width = cw

# 헤더 행
for ci, h_txt in enumerate(headers):
    stylecell(tbl2.cell(0, ci), h_txt,
              size=Pt(10.5), bold=True, color=WHITE, bg=NAVY)

# 데이터 행
for ri, (name, init, curr, new, dept, price, changed) in enumerate(full):
    row_idx = ri + 1
    base_bg = WHITE if ri % 2 == 0 else BLUE_LIGHT
    changed_bg = RGBColor(0xFF, 0xF0, 0xF0)

    for ci, val in enumerate([name, init, curr, new, dept, price]):
        c = tbl2.cell(row_idx, ci)
        if changed:
            if ci == 0:
                stylecell(c, val, size=Pt(9.5), bold=True, color=RED,
                          bg=changed_bg, align=PP_ALIGN.CENTER)
                stylecell(tbl2.cell(row_idx, 0), "★ " + val,
                          size=Pt(9.5), bold=True, color=RED,
                          bg=changed_bg, align=PP_ALIGN.CENTER)
            elif ci == 2:
                stylecell(c, val, size=Pt(9), color=RED, bg=changed_bg, italic=True)
            elif ci == 3:
                stylecell(c, val, size=Pt(9), bold=True, color=GREEN, bg=GREEN_LIGHT)
            else:
                stylecell(c, val, size=Pt(9), color=GRAY_DARK, bg=changed_bg)
        else:
            stylecell(c, val, size=Pt(9), color=(NAVY if ci==0 else GRAY_DARK),
                      bold=(ci==0), bg=base_bg)

txbox(slide, "★  붉은 항목 = 이번 개정 대상 3종  |  나머지 항목 현행 유지",
      Inches(0.35), Inches(7.18), Inches(10), Inches(0.28),
      size=Pt(9), color=GRAY_MID, italic=True)


# ════════════════════════════════════════
# SLIDE 6 — 기대 효과
# ════════════════════════════════════════
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, W, H, fill=GRAY_BG)
header_bar(slide, "04  기대 효과", "6 / 7")

effects = [
    (RED,    "💰",  "재경 비용 절감",
     "연간 사원증 제작비·부속품 구입비 절감",
     [
         "사원증 재발급 단가: 46,750원/건",
         "케이스·목걸이 포함 세트 비용 절감",
         "총무팀 발주·행정 처리 공수 감소",
         "소모성 비용 누수 방지 효과",
     ]),
    (NAVY,   "🔒",  "출입 보안 강화",
     "분실 경각심 제고 → 보안 사고 예방",
     [
         "분실 즉시 본인 책임 인식 강화",
         "병원·재단 내 비인가 출입 위험 차단",
         "분실 신고 즉시성 향상 기대",
         "보안 사고 사전 예방 효과",
     ]),
    (GREEN,  "✅",  "직원 불만 최소화",
     "합리적 예외 조항으로 수용성 확보",
     [
         "근속 5년 이상 자연 노후화 1회 무상 교체",
         "제도 변경 반발 심리 사전 차단",
         "최초 지급 기준 변경 없이 유지",
         "유예기간 1주일 부여로 적응 지원",
     ]),
]

for i, (accent, icon, title, subtitle, bullets) in enumerate(effects):
    lft = Inches(0.35) + i * Inches(4.3)
    top = Inches(0.95)
    bw  = Inches(4.1)
    bh  = Inches(6.1)
    rect(slide, lft, top, bw, bh, fill=WHITE, line=GRAY_LINE, lw=Pt(0.8))
    rect(slide, lft, top, bw, Inches(0.6), fill=accent)
    txbox(slide, f"{icon}  {title}",
          lft + Inches(0.15), top + Inches(0.09), bw - Inches(0.3), Inches(0.44),
          size=Pt(13), bold=True, color=WHITE)
    txbox(slide, subtitle,
          lft + Inches(0.15), top + Inches(0.72), bw - Inches(0.3), Inches(0.38),
          size=Pt(11), bold=True, color=accent)
    rect(slide, lft + Inches(0.15), top + Inches(1.12),
         bw - Inches(0.3), Inches(0.02), fill=GRAY_LINE)
    for j, b in enumerate(bullets):
        rect(slide, lft + Inches(0.18), top + Inches(1.25) + j * Inches(1.1),
             Inches(0.05), Inches(0.05), fill=accent)
        txbox(slide, b,
              lft + Inches(0.32), top + Inches(1.18) + j * Inches(1.1),
              bw - Inches(0.5), Inches(1.0),
              size=Pt(10.5), color=GRAY_DARK, wrap=True)


# ════════════════════════════════════════
# SLIDE 7 — 향후 일정 및 결재 요청
# ════════════════════════════════════════
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, W, H, fill=GRAY_BG)
header_bar(slide, "05  향후 일정 및 결재 요청", "7 / 7")

# 타임라인
steps = [
    (BLUE,  "STEP 1",  "2026. 05",     "CEO 결재",
     "• 본 개선(안) CEO 승인\n• 총무팀 내부 검토 완료"),
    (TEAL,  "STEP 2",  "2026. 05",     "사내 규정 개정",
     "• 복리후생 지급 기준 규정 개정\n• 재경팀 / 인사팀 프로세스 연동"),
    (NAVY,  "STEP 3",  "2026. 05",     "전사 공지",
     "• 전 직원 공지 및 전파\n• 조직문화 유예기간 1주일 부여"),
    (RED,   "STEP 4",  "2026. 06. 01","변경 기준 시행",
     "• 변경된 지급 기준 전격 시행\n• 총무팀 운영 현황 모니터링"),
]

for i, (accent, step, date, title, body) in enumerate(steps):
    lft = Inches(0.35) + i * Inches(3.2)
    # 연결선
    if i < len(steps) - 1:
        rect(slide, lft + Inches(3.0), Inches(2.18),
             Inches(0.4), Inches(0.08), fill=GRAY_LINE)
        txbox(slide, "▶",
              lft + Inches(2.98), Inches(1.98), Inches(0.44), Inches(0.4),
              size=Pt(16), color=GRAY_LINE, align=PP_ALIGN.CENTER)

    bw = Inches(3.0)
    # 스텝 뱃지
    rect(slide, lft, Inches(1.0), bw, Inches(0.45), fill=accent)
    txbox(slide, step,
          lft, Inches(1.05), bw, Inches(0.35),
          size=Pt(11), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # 날짜
    rect(slide, lft, Inches(1.48), bw, Inches(0.52), fill=BLUE_LIGHT, line=accent, lw=Pt(1.0))
    txbox(slide, date,
          lft, Inches(1.52), bw, Inches(0.42),
          size=Pt(14), bold=True, color=accent, align=PP_ALIGN.CENTER)
    # 제목
    rect(slide, lft, Inches(2.05), bw, Inches(0.52), fill=WHITE, line=GRAY_LINE, lw=Pt(0.5))
    txbox(slide, title,
          lft + Inches(0.1), Inches(2.1), bw - Inches(0.2), Inches(0.42),
          size=Pt(12), bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    # 본문
    rect(slide, lft, Inches(2.6), bw, Inches(2.2),
         fill=WHITE, line=GRAY_LINE, lw=Pt(0.5))
    txbox(slide, body,
          lft + Inches(0.12), Inches(2.7), bw - Inches(0.25), Inches(2.0),
          size=Pt(10.5), color=GRAY_DARK, wrap=True)

# 결재 요청 박스
rect(slide, Inches(0.35), Inches(5.0), W - Inches(0.7), Inches(2.1),
     fill=NAVY, line=GOLD, lw=Pt(2.0))
txbox(slide, "■  결재 요청 사항",
      Inches(0.6), Inches(5.1), Inches(4), Inches(0.38),
      size=Pt(14), bold=True, color=GOLD)
txbox(slide,
      "신규입사자 지원물품 중 사원증·사원증 케이스·사원증 목걸이(3종)에 대한 재발급 기준을\n"
      "「분실·파손 시 급여 100% 공제 후 재지급」으로 변경하는 안에 대해 CEO 결재를 요청드립니다.\n"
      "근속 5년 이상 자연 노후화의 경우 1회 무상 교체 예외 조항을 포함합니다.",
      Inches(0.6), Inches(5.52), W - Inches(1.2), Inches(1.45),
      size=Pt(11.5), color=WHITE, wrap=True)

# 결재 서명란
sign_labels = ["기안", "검토", "승인"]
for i, lbl in enumerate(sign_labels):
    sx = Inches(9.5) + i * Inches(1.2)
    sy = Inches(5.08)
    rect(slide, sx, sy, Inches(1.1), Inches(0.9),
         fill=RGBColor(0x10, 0x35, 0x65), line=GOLD, lw=Pt(0.8))
    txbox(slide, lbl,
          sx, sy + Inches(0.02), Inches(1.1), Inches(0.28),
          size=Pt(9), bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    txbox(slide, "",
          sx, sy + Inches(0.3), Inches(1.1), Inches(0.55),
          size=Pt(9), color=WHITE, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════
# 저장
# ════════════════════════════════════════
out = r"C:\Users\최다빈\Desktop\PPT\신규입사자_지원물품_CEO보고서_v2.pptx"
prs.save(out)
print("완료:", out)
